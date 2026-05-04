"""Generate a DOCX report with the scientific contributions of CityLearn v3 MADRL."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOCS_DIR = PROJECT_ROOT / "docs"
OUTPUT_DOCX = DOCS_DIR / "APORTES_CIENTIFICOS_CITYLEARN_V3_MADRL.docx"


TITLE = (
    "Aportes científicos del proyecto CityLearn v3 MADRL para gestión "
    "coordinada de flexibilidad energética, emisiones de carbono y eficiencia "
    "económica en comunidades inteligentes"
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells

    for cell, header in zip(header_cells, headers):
        set_cell_shading(cell, "1F4E79")
        set_cell_text(cell, header, bold=True, color="FFFFFF")

    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)

    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(18)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.name = "Aptos"
    styles["Heading 3"].font.size = Pt(11.5)


def add_cover(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)

    document.add_paragraph()
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Documento técnico de aportes a la ciencia").bold = True

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    document.add_paragraph()
    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.add_run(
        "Repositorio: d:\\MADRLCitytleranflexresdr\n"
        "Base experimental: CityLearn v2 + capa CityLearn v3 MADRL\n"
        "Dataset: citylearn_challenge_2022_phase_all_plus_evs"
    )

    document.add_section(WD_SECTION.NEW_PAGE)


def build_document() -> Document:
    document = Document()
    configure_document(document)
    add_cover(document)

    document.add_heading("1. Síntesis ejecutiva", level=1)
    document.add_paragraph(
        "El proyecto aporta una extensión reproducible de CityLearn v2 hacia una "
        "capa CityLearn v3 orientada a aprendizaje por refuerzo profundo "
        "multiagente. La contribución científica central consiste en convertir "
        "una comunidad de 17 edificios con EV/V2G, PV, baterías, precios y "
        "carbono en un banco experimental MADRL bajo Dec-POMDP y CTDE, sin "
        "abandonar los KPIs oficiales de CityLearn v2."
    )
    document.add_paragraph(
        "El resultado no es solo una integración de código: es una metodología "
        "comparativa para estudiar flexibilidad energética, emisiones de CO2 y "
        "costos económicos con cuatro familias de algoritmos MADRL conectadas "
        "al mismo entorno, al mismo dataset y al mismo protocolo de artefactos."
    )

    document.add_heading("2. Aporte científico general", level=1)
    document.add_paragraph(
        "El aporte general del proyecto es un marco experimental abierto y "
        "trazable para evaluar control coordinado de comunidades inteligentes "
        "mediante MADRL, preservando la validez comparativa de CityLearn v2 y "
        "añadiendo capacidades de entrenamiento descentralizado multiagente "
        "con críticos o estados centralizados durante entrenamiento."
    )

    document.add_heading("3. Aportes científicos principales", level=1)
    add_table(
        document,
        ["N.º", "Aporte", "Descripción científica", "Evidencia implementada"],
        [
            [
                "1",
                "Extensión CityLearn v2 -> CityLearn v3 MADRL",
                "Propone una capa experimental que conserva CityLearn v2 como simulador y lo adapta a entrenamiento MADRL profundo.",
                "`CityLearn/citylearn/v3`, `citylearn_v3_training_common.py`, scripts `train_citylearn_v3_*.py`.",
            ],
            [
                "2",
                "Formulación Dec-POMDP para comunidades de edificios",
                "Cada edificio se modela como agente con observación local, acción local y coordinación distrital bajo información parcial.",
                "`CityLearn/citylearn/dec_pomdp.py`, adaptadores PettingZoo-like y wrappers MADRL.",
            ],
            [
                "3",
                "Contrato CTDE común",
                "Estandariza entrenamiento centralizado y ejecución descentralizada para algoritmos heterogéneos.",
                "`CityLearnHARLEnv`, `CityLearnSMACDiscreteEnv`, `CityLearnOffPolicyVecEnv`, `CityLearnMAACVecEnv`.",
            ],
            [
                "4",
                "Comparación de cuatro MADRL oficiales",
                "Integra HAPPO, MASAC, MATD3 y MAAC sobre el mismo dataset, horizonte y KPIs, reduciendo sesgo de entorno.",
                "`external/HARL`, `external/MARL`, `external/off-policy`, `external/MAAC`.",
            ],
            [
                "5",
                "Marco multiobjetivo de tres ejes",
                "Organiza la evaluación científica en flexibilidad, emisiones de CO2 y costos energéticos.",
                "`CityLearn/citylearn/v3/objectives.py`, `scenario_manager.py`.",
            ],
            [
                "6",
                "Uso de KPIs CityLearn v2 sin confundirlos con métricas del proyecto",
                "Distingue los ejes científicos OE1/OE2/OE3 de los KPIs técnicos de CityLearn v2.",
                "`objective_manifest()`, `evaluate_objectives()`, `objective_kpis.csv`.",
            ],
            [
                "7",
                "Protocolo de artefactos reproducibles",
                "Cada corrida produce JSON, CSV, checkpoints, figuras, tablas y progreso vivo con trazabilidad por algoritmo/eje/seed.",
                "`results.json`, `training_summary.json`, `timeseries.csv`, `trace.csv`, `checkpoint_manifest.json`.",
            ],
            [
                "8",
                "Benchmark CityLearn v2 vs CityLearn v3",
                "Permite validar si la capa MADRL v3 mejora o caracteriza mejor las soluciones frente a agentes originales v2.",
                "`benchmark_citylearn_v2_agents.py`, `compare_citylearn_v2_vs_v3_madrl.py`.",
            ],
            [
                "9",
                "Monitoreo técnico en vivo",
                "Hace observable el entrenamiento con GPU, rewards, costo, CO2, carga neta y estado por job.",
                "`monitor_citylearn_v3_official_training.ps1`, `live_progress.json`.",
            ],
            [
                "10",
                "Escalabilidad experimental",
                "El protocolo `-Scenario ALL` ejecuta 12 corridas base: tres ejes por cuatro MADRL.",
                "`launch_citylearn_v3_official_training.ps1 -Scenario ALL`.",
            ],
        ],
    )

    document.add_heading("4. Aportes por eje científico", level=1)
    add_table(
        document,
        ["Eje", "Objetivo científico", "KPIs y evidencia", "Aporte esperado"],
        [
            [
                "OE1 - Flexibilidad energética",
                "Aumentar la capacidad de desplazar cargas y aprovechar almacenamiento, EVs y autoconsumo.",
                "`peak_average`, `ramping_average`, `one_minus_load_factor_average`, KPIs PV, batería y EV.",
                "Permite estudiar políticas colaborativas para reducir picos, suavizar rampas y coordinar activos flexibles.",
            ],
            [
                "OE2 - Emisiones de CO2",
                "Reducir la huella ambiental del distrito, minimizando importaciones en horas de alta intensidad de carbono.",
                "`carbon_emissions`, `carbon_emissions_control`, `carbon_emissions_baseline`, `carbon_emissions_delta`.",
                "Introduce evaluación carbon-aware con trazabilidad entre operación eléctrica y emisiones de gases de efecto invernadero.",
            ],
            [
                "OE3 - Costos energéticos",
                "Optimizar el gasto energético, reduciendo picos de demanda y aprovechando tarifas dinámicas.",
                "`electricity_cost`, `electricity_cost_delta`, `price_signal_deviation`, KPIs de pico y rampa de costo.",
                "Permite analizar si la coordinación MADRL reduce costos sin sacrificar flexibilidad ni desempeño ambiental.",
            ],
        ],
    )

    document.add_heading("5. Aportes metodológicos", level=1)
    add_bullets(
        document,
        [
            "Diseño de una matriz experimental común: 4 algoritmos MADRL x 3 ejes x semillas reproducibles.",
            "Separación entre objetivos científicos del proyecto y KPIs técnicos de CityLearn v2.",
            "Comparación controlada contra agentes originales CityLearn v2 usando el mismo dataset.",
            "Registro de series temporales distritales y trazas por agente para análisis macro y micro.",
            "Uso de salidas estándar que permiten auditoría posterior: JSON, CSV, PNG y checkpoints.",
            "Ejecución secuencial GPU-safe para equipos con memoria limitada, sin alterar el protocolo científico.",
        ],
    )

    document.add_heading("6. Aportes algorítmicos y de integración", level=1)
    add_table(
        document,
        ["Algoritmo", "Backend activo", "Aporte al proyecto"],
        [
            [
                "HAPPO",
                "`external/HARL`",
                "Introduce un actor-critic on-policy multiagente con entrenamiento centralizado y políticas por agente.",
            ],
            [
                "MASAC",
                "`external/MARL/src`",
                "Permite explorar aprendizaje con entropía y acciones discretizadas sobre un estado global tipo SMAC.",
            ],
            [
                "MATD3",
                "`external/off-policy`",
                "Permite control continuo off-policy con replay buffer, actores por agente y críticos centralizados.",
            ],
            [
                "MAAC",
                "`external/MAAC`",
                "Agrega coordinación mediante crítico con atención para estudiar influencia entre agentes.",
            ],
        ],
    )

    document.add_heading("7. Aportes de software científico", level=1)
    add_bullets(
        document,
        [
            "Una capa `CityLearn/citylearn/v3` separada del núcleo CityLearn v2, lo que reduce acoplamiento y facilita evolución futura.",
            "Scripts de entrenamiento independientes por MADRL, evitando doble implementación de algoritmos dentro de `citylearn.agents`.",
            "Manifiesto de backends con rutas, repositorios y commits para reproducibilidad.",
            "Generadores de documentación: notebook tutorial, planos PDF/PNG y Markdown renderizable.",
            "Celda de Google Colab GPU para ejecutar los mismos MADRL en A100/T4/V100 sin depender de PowerShell.",
        ],
    )

    document.add_heading("8. Aportes de reproducibilidad y trazabilidad", level=1)
    add_numbered(
        document,
        [
            "Cada corrida queda identificada por algoritmo, escenario, seed y carpeta de salida.",
            "Los hiperparámetros quedan en `training_summary.json`.",
            "Los checkpoints se referencian mediante `checkpoint_manifest.json`.",
            "Las decisiones por agente quedan en `trace.csv`.",
            "El desempeño distrital queda en `timeseries.csv`.",
            "Las figuras y tablas se regeneran desde los artefactos guardados.",
            "El comparador v2 vs v3 separa la evaluación científica del entrenamiento.",
        ],
    )

    document.add_heading("9. Aportes aplicados al sistema energético", level=1)
    document.add_paragraph(
        "El proyecto contribuye a la gestión de comunidades grid-interactive al "
        "proporcionar un protocolo para coordinar baterías, EVs, PV y consumo "
        "edilicio bajo señales de precio y carbono. Esto habilita estudios "
        "sobre desplazamiento de demanda, autoconsumo, V2G, reducción de "
        "emisiones y eficiencia económica con un mismo entorno experimental."
    )

    document.add_heading("10. Aportes para tesis y futuras investigaciones", level=1)
    add_bullets(
        document,
        [
            "Base para tesis de maestría centrada en control MADRL multiobjetivo de comunidades inteligentes.",
            "Plataforma para estudiar sensibilidad por semillas, penetración EV, disponibilidad PV y escenarios tarifarios.",
            "Punto de partida para comparar nuevos algoritmos MADRL manteniendo el mismo contrato de entorno.",
            "Posibilidad de extender el análisis hacia equidad entre edificios, curvas de Pareto y robustez ante fallas.",
            "Material pedagógico reproducible mediante notebook comentado y diagramas renderizables.",
        ],
    )

    document.add_heading("11. Evidencia implementada en el repositorio", level=1)
    add_table(
        document,
        ["Evidencia", "Ruta"],
        [
            ["Notebook tutorial comentado", "`CityLearn/examples/madrl_citylearn_v3_tutorial.ipynb`"],
            ["Arquitectura renderizable", "`docs/ARQUITECTURA_Y_FLUJO_TRABAJO_CITYLEARN_V3_MADRL.md`"],
            ["Plano real implementado", "`docs/PLANO_REAL_IMPLEMENTADO_CITYLEARN_V3_MADRL.pdf`"],
            ["Launcher oficial", "`CityLearn/scripts/launch_citylearn_v3_official_training.ps1`"],
            ["Monitor vivo", "`CityLearn/scripts/monitor_citylearn_v3_official_training.ps1`"],
            ["Entrenadores MADRL", "`CityLearn/scripts/train_citylearn_v3_*.py`"],
            ["Benchmark v2", "`CityLearn/scripts/benchmark_citylearn_v2_agents.py`"],
            ["Comparador v2 vs v3", "`CityLearn/scripts/compare_citylearn_v2_vs_v3_madrl.py`"],
        ],
    )

    document.add_heading("12. Alcance y cautelas científicas", level=1)
    document.add_paragraph(
        "Los aportes anteriores describen la arquitectura, metodología y "
        "capacidad experimental implementada. La magnitud empírica de la mejora "
        "de CityLearn v3 MADRL sobre CityLearn v2 debe establecerse con los "
        "resultados completos de entrenamiento, benchmark y comparación final. "
        "Por rigor, el proyecto debe reportar no solo la mejor corrida, sino "
        "también estabilidad por semilla, costo computacional, limitaciones de "
        "discretización en algunos backends y sensibilidad de los pesos "
        "multiobjetivo."
    )

    document.add_heading("13. Conclusión", level=1)
    document.add_paragraph(
        "El proyecto aporta a la ciencia un banco experimental MADRL reproducible "
        "para comunidades energéticas inteligentes, articulado sobre CityLearn v2 "
        "y extendido hacia CityLearn v3 con Dec-POMDP, CTDE, cuatro backends "
        "MADRL oficiales, tres ejes científicos y comparación directa contra "
        "baseline. Su principal valor científico está en la integración rigurosa, "
        "la trazabilidad de resultados y la posibilidad de evaluar políticas "
        "colaborativas bajo flexibilidad, carbono y costo en un mismo marco."
    )

    document.add_heading("14. Referencias base", level=1)
    references = [
        "Nweye, K., Kaspar, K., Buscemi, G., Fonseca, T., Pinto, G., Ghose, D., Duddukuru, S., Pratapa, P., Li, H., Mohammadi, J., Lino Ferreira, L., Hong, T., Ouf, M., Capozzoli, A., & Nagy, Z. (2025). CityLearn v2: Energy-flexible, resilient, occupant-centric, and carbon-aware management of grid-interactive communities. Journal of Building Performance Simulation, 18(1), 17-38. https://doi.org/10.1080/19401493.2024.2418813",
        "Vázquez-Canteli, J. R., Dey, S., Henze, G., & Nagy, Z. (2020). CityLearn: Standardizing research in multi-agent reinforcement learning for demand response and urban energy management. arXiv. https://doi.org/10.48550/arXiv.2012.10504",
        "Zhong, Y., Kuba, J. G., Feng, X., Hu, S., Ji, J., & Yang, Y. (2024). Heterogeneous-agent reinforcement learning. Journal of Machine Learning Research, 25(32), 1-67.",
        "Iqbal, S., & Sha, F. (2019). Actor-attention-critic for multi-agent reinforcement learning. arXiv. https://doi.org/10.48550/arXiv.1810.02912",
        "Ackermann, J., Gabler, V., Osa, T., & Sugiyama, M. (2019). Reducing overestimation bias in multi-agent domains using double centralized critics. arXiv. https://doi.org/10.48550/arXiv.1910.01465",
        "Pu, Y., Wang, S., Yang, R., Yao, X., & Li, B. (2021). Decomposed soft actor-critic method for cooperative multi-agent reinforcement learning. arXiv. https://doi.org/10.48550/arXiv.2104.06655",
    ]
    for reference in references:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.7)
        paragraph.paragraph_format.first_line_indent = Cm(-0.7)
        paragraph.add_run(reference)

    return document


def main() -> int:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
